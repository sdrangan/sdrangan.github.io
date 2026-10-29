"""Write the NOFO-1 and NOFO-2 product spreadsheets for one report period.

Reads pis.yml (awards and report period), candidates.yml (from scrape.py),
decisions.yml and products.yml, and writes one .xlsx per award to output/.
Prints the papers in the period that are not yet in decisions.yml.

Usage:  python _pubs/build.py
"""

import re
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "output"

AWARD_CHOICES = {"nofo1": ["nofo1"], "nofo2": ["nofo2"], "both": ["nofo1", "nofo2"]}
OVERRIDES = ("title", "authors", "venue", "type", "url", "date")
COLUMNS = ["Grant Recipient", "Performance Report", "Product Title", "Product Type",
           "URL", "Publication Name", "Authors"]
WIDTHS = [32, 44, 70, 14, 45, 45, 60]


def load(name):
    return yaml.safe_load((HERE / name).read_text(encoding="utf-8")) or {}


def norm_title(title):
    return re.sub(r"[^a-z0-9]", "", title.lower())


def resolve(entry, decision):
    """Apply a decision to a scraped entry; returns None if excluded."""
    e = dict(entry)
    e["date"] = str(e["date"])
    if decision is None:
        e["review"] = True
        return e
    award = decision.get("award", "default")
    if award == "exclude":
        return None
    if award != "default":
        if award not in AWARD_CHOICES:
            raise ValueError(f"{entry['id']}: unknown award '{award}'")
        e["awards"] = AWARD_CHOICES[award]
    for k in OVERRIDES:
        if k in decision:
            e[k] = str(decision[k]) if k == "date" else decision[k]
    if "report" in decision:
        e["report"] = str(decision["report"])
    return e


def in_report(e, start, end):
    if e.get("report"):
        return e["report"] == end
    return start <= e["date"] <= end


def main():
    cfg = load("pis.yml")
    awards = cfg["awards"]
    start, end = str(cfg["report"]["start"]), str(cfg["report"]["end"])
    candidates = load("candidates.yml") or []
    decisions = load("decisions.yml")
    products = load("products.yml") or []

    known = {e["id"] for e in candidates}
    for key in decisions:
        if key not in known:
            print(f"warning: decisions.yml lists {key}, which is not in candidates.yml")

    entries = [e for e in (resolve(c, decisions.get(c["id"])) for c in candidates) if e]
    entries = [e for e in entries if in_report(e, start, end)]

    # A preprint whose published version is in the same report is left out.
    # OpenAlex links some pairs (related); others only share a title.
    ids = {e["id"] for e in entries}
    published = {norm_title(e["title"]) for e in entries if e["type"] != "Preprint"}
    entries = [e for e in entries
               if not (e["type"] == "Preprint"
                       and (any(r in ids for r in e.get("related", []))
                            or norm_title(e["title"]) in published))]

    for p in products:
        missing = [k for k in ("title", "type", "date", "award") if k not in p]
        if missing:
            raise ValueError(f"products.yml: '{p.get('title', '?')}' is missing {', '.join(missing)}")
        if p["award"] not in AWARD_CHOICES:
            raise ValueError(f"products.yml: '{p['title']}' has unknown award '{p['award']}'")
        e = {"id": "product", "title": p["title"], "type": p["type"], "date": str(p["date"]),
             "awards": AWARD_CHOICES[p["award"]], "venue": p.get("venue", ""),
             "authors": p.get("authors", ""), "url": p.get("url", "")}
        if "report" in p:
            e["report"] = str(p["report"])
        if in_report(e, start, end):
            entries.append(e)

    entries.sort(key=lambda e: (e["date"], e["title"]))
    y, m, d = (int(x) for x in end.split("-"))
    OUT_DIR.mkdir(exist_ok=True)

    for key, award in awards.items():
        rows = [e for e in entries if key in e["awards"]]
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.append(COLUMNS)
        for e in rows:
            ws.append([award["recipient"], f"{award['recipient']} ({y} {m}/{d} Report)",
                       e["title"], e["type"], e["url"], e["venue"], e["authors"]])
        for i, width in enumerate(WIDTHS):
            ws.column_dimensions[chr(65 + i)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        path = OUT_DIR / f"{award['file']} {end}.xlsx"
        wb.save(path)
        print(f"{len(rows):3d} items -> {path.relative_to(HERE)}")

    review = [e for e in entries if e.get("review")]
    if review:
        print(f"\n{len(review)} items in the period are not in decisions.yml yet "
              "(included above with their default awards):")
        for e in review:
            print(f"  {e['date']}  {'+'.join(e['awards']):11s} {e['id']}\n      {e['title'][:90]}")


if __name__ == "__main__":
    main()
